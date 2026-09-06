import re
from io import BytesIO

SECTION_TITLES = {
    "project_plan.md": "Project Plan",
    "requirements.md": "Requirements",
    "architecture.md": "Architecture",
    "plan.md": "Planning",
    "api_spec.md": "Backend API Specification",
    "ui_plan.md": "Frontend UI Plan",
    "testing.md": "QA & Testing",
    "README.md": "Documentation",
}

SECTION_ORDER = [
    "requirements.md",
    "architecture.md",
    "plan.md",
    "api_spec.md",
    "ui_plan.md",
    "testing.md",
    "README.md",
    "project_plan.md",
]


def split_report_markdown(markdown: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    current_file: str | None = None
    current_lines: list[str] = []
    for line in markdown.splitlines():
        marker = re.match(r"<!-- section: (.+?) -->", line)
        if marker:
            if current_file is not None:
                sections.append((current_file, "\n".join(current_lines).strip()))
            current_file = marker.group(1)
            current_lines = []
        else:
            current_lines.append(line)
    if current_file is not None:
        sections.append((current_file, "\n".join(current_lines).strip()))
    return sections


def md_to_blocks(text: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    in_code = False
    code_lines: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                blocks.append(("code", "\n".join(code_lines)))
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            blocks.append((f"h{level}", heading.group(2)))
        elif re.match(r"^[-*]\s+", stripped):
            blocks.append(("bullet", re.sub(r"^[-*]\s+", "", stripped)))
        elif re.match(r"^\d+\.\s+", stripped):
            blocks.append(("numbered", re.sub(r"^\d+\.\s+", "", stripped)))
        elif stripped.startswith("|"):
            blocks.append(("code", stripped))
        elif stripped.startswith("!"):
            blocks.append(("code", stripped))
        else:
            blocks.append(("text", stripped))
    if in_code and code_lines:
        blocks.append(("code", "\n".join(code_lines)))
    return blocks


def _strip_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    return text


def _to_pdf_rich(text: str) -> str:
    text = re.sub(r"&", "&amp;", text)
    text = re.sub(r"<", "&lt;", text)
    text = re.sub(r">", "&gt;", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r"<font name='Courier'>\1</font>", text)
    return text


def _ordered_sections(sections: list[tuple[str, str]]) -> list[tuple[str, str]]:
    by_file = {filename: body for filename, body in sections}
    ordered = [(f, by_file[f]) for f in SECTION_ORDER if f in by_file]
    for filename, body in sections:
        if filename not in SECTION_ORDER:
            ordered.append((filename, body))
    return ordered


def build_pdf_bytes(
    project_title: str,
    report_markdown: str,
    metrics: dict,
    agent_runs: list[dict],
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"AutoDev Report - {project_title}",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", fontName="Helvetica-Bold", fontSize=30, leading=36))
    styles.add(ParagraphStyle(name="CoverSub", fontName="Helvetica", fontSize=14, leading=20, textColor=colors.HexColor("#555555")))
    styles.add(ParagraphStyle(name="SectionTitle", fontName="Helvetica-Bold", fontSize=16, leading=20, spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="SectionSub", fontName="Helvetica-Bold", fontSize=12, leading=16, spaceBefore=6, spaceAfter=2))
    styles.add(ParagraphStyle(name="Body", fontName="Helvetica", fontSize=10, leading=14, spaceAfter=4))

    story: list = []

    story.append(Spacer(1, 60 * mm))
    story.append(Paragraph(_to_pdf_rich(project_title), styles["CoverTitle"]))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("AutoDev Project Report", styles["CoverSub"]))
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph("Generated by the AutoDev multi-agent pipeline", styles["CoverSub"]))
    story.append(PageBreak())

    story.append(Paragraph("Executive Summary", styles["SectionTitle"]))
    summary_rows = [
        ["Project", project_title],
        ["Quality score", f"{metrics.get('quality', '-')}"],
        ["Execution time", f"{metrics.get('execution_time_ms', '-')} ms"],
        ["Estimated cost", f"${metrics.get('cost', '-')}"],
        [
            "Tokens",
            f"{metrics.get('tokens', {}).get('input', 0)} in / "
            f"{metrics.get('tokens', {}).get('output', 0)} out",
        ],
    ]
    summary_table = Table([[Paragraph(k, styles["Body"]), Paragraph(v, styles["Body"])] for k, v in summary_rows], colWidths=[40 * mm, 110 * mm])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F4F6F8")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D2DA")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(summary_table)
    story.append(PageBreak())

    sections = _ordered_sections(split_report_markdown(report_markdown))
    for filename, body in sections:
        story.append(Paragraph(SECTION_TITLES.get(filename, filename), styles["SectionTitle"]))
        for kind, content in md_to_blocks(body):
            if kind == "code":
                story.append(Preformatted(content, ParagraphStyle(name="Code", fontName="Courier", fontSize=8, leading=10, leftIndent=6 * mm)))
            elif kind.startswith("h"):
                story.append(Paragraph(_to_pdf_rich(content), styles["SectionSub"]))
            elif kind in ("bullet", "numbered"):
                story.append(Paragraph(f"&bull; {_to_pdf_rich(content)}", styles["Body"]))
            else:
                story.append(Paragraph(_to_pdf_rich(content), styles["Body"]))
        story.append(Spacer(1, 4 * mm))
    story.append(PageBreak())

    story.append(Paragraph("Metrics", styles["SectionTitle"]))
    metric_rows = [
        ["Quality score", str(metrics.get("quality", "-"))],
        ["Execution time", f"{metrics.get('execution_time_ms', '-')} ms"],
        ["Estimated cost", f"${metrics.get('cost', '-')}"],
        ["Input tokens", str(metrics.get("tokens", {}).get("input", 0))],
        ["Output tokens", str(metrics.get("tokens", {}).get("output", 0))],
    ]
    metric_table = Table([[Paragraph(k, styles["Body"]), Paragraph(v, styles["Body"])] for k, v in metric_rows], colWidths=[40 * mm, 110 * mm])
    metric_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D2DA")),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(metric_table)

    if agent_runs:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("Appendix: Agent Runs", styles["SectionTitle"]))
        header = ["Agent", "Status", "Confidence", "Duration (ms)", "Summary"]
        rows = [header] + [
            [Paragraph(_to_pdf_rich(str(run.get("agent_name", ""))), styles["Body"]),
             Paragraph(_to_pdf_rich(str(run.get("status", ""))), styles["Body"]),
             Paragraph(_to_pdf_rich(str(run.get("confidence", ""))), styles["Body"]),
             Paragraph(_to_pdf_rich(str(run.get("execution_ms", ""))), styles["Body"]),
             Paragraph(_to_pdf_rich(str(run.get("summary", ""))), styles["Body"])]
            for run in agent_runs
        ]
        run_table = Table(rows, colWidths=[28 * mm, 18 * mm, 24 * mm, 26 * mm, 54 * mm])
        run_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EBF0")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D2DA")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(run_table)

    doc.build(story)
    return buffer.getvalue()


def build_docx_bytes(
    project_title: str,
    report_markdown: str,
    metrics: dict,
    agent_runs: list[dict],
) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_heading("AutoDev Project Report", 0)
    document.add_heading("Executive Summary", 1)
    summary = document.add_table(rows=0, cols=2)
    summary.style = "Light Grid Accent 1"
    summary_data = [
        ("Project", project_title),
        ("Quality score", str(metrics.get("quality", "-"))),
        ("Execution time", f"{metrics.get('execution_time_ms', '-')} ms"),
        ("Estimated cost", f"${metrics.get('cost', '-')}"),
        (
            "Tokens",
            f"{metrics.get('tokens', {}).get('input', 0)} in / "
            f"{metrics.get('tokens', {}).get('output', 0)} out",
        ),
    ]
    for key, value in summary_data:
        cells = summary.add_row().cells
        cells[0].text = key
        cells[1].text = value

    for filename, body in _ordered_sections(split_report_markdown(report_markdown)):
        document.add_heading(SECTION_TITLES.get(filename, filename), 1)
        for kind, content in md_to_blocks(body):
            if kind == "code":
                paragraph = document.add_paragraph()
                run = paragraph.add_run(content)
                run.font.name = "Courier New"
                run.font.size = Pt(8)
            elif kind.startswith("h"):
                document.add_heading(content, int(kind[1]))
            elif kind == "bullet":
                document.add_paragraph(content, style="List Bullet")
            elif kind == "numbered":
                document.add_paragraph(content, style="List Number")
            else:
                document.add_paragraph(content)

    document.add_heading("Metrics", 1)
    metrics_table = document.add_table(rows=0, cols=2)
    metrics_table.style = "Light Grid Accent 1"
    for key, value in [
        ("Quality score", str(metrics.get("quality", "-"))),
        ("Execution time", f"{metrics.get('execution_time_ms', '-')} ms"),
        ("Estimated cost", f"${metrics.get('cost', '-')}"),
        ("Input tokens", str(metrics.get("tokens", {}).get("input", 0))),
        ("Output tokens", str(metrics.get("tokens", {}).get("output", 0))),
    ]:
        cells = metrics_table.add_row().cells
        cells[0].text = key
        cells[1].text = value

    if agent_runs:
        document.add_heading("Appendix: Agent Runs", 1)
        runs_table = document.add_table(rows=1, cols=5)
        runs_table.style = "Light Grid Accent 1"
        for index, header in enumerate(["Agent", "Status", "Confidence", "Duration (ms)", "Summary"]):
            runs_table.rows[0].cells[index].text = header
        for run in agent_runs:
            cells = runs_table.add_row().cells
            cells[0].text = str(run.get("agent_name", ""))
            cells[1].text = str(run.get("status", ""))
            cells[2].text = str(run.get("confidence", ""))
            cells[3].text = str(run.get("execution_ms", ""))
            cells[4].text = str(run.get("summary", ""))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
